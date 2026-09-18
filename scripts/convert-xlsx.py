"""Convert the DND site-list spreadsheet into src/data/sites.json."""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import openpyxl


def clean_str(value) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value).strip())
    return text or None


def clean_num(value):
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number == int(number):
        return int(number)
    return number


def convert(xlsx_path: Path) -> dict:
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet = workbook["DND"]
    sites = []

    for row in sheet.iter_rows(min_row=4, max_row=sheet.max_row):
        serial = row[1].value
        site_code = clean_str(row[3].value)
        if serial is None or not site_code:
            continue

        latitude = clean_num(row[12].value)
        longitude = clean_num(row[13].value)
        if latitude is None or longitude is None:
            continue

        width = clean_num(row[6].value)
        height = clean_num(row[7].value)
        qty = clean_num(row[8].value)
        area = clean_num(row[9].value)
        if area is None and width is not None and height is not None:
            area = width * height * (qty if qty is not None else 1)

        # Spreadsheet formula: =HYPERLINK("http://maps.google.com/?q="&M&","&N,"Click here to map")
        google_maps_url = f"http://maps.google.com/?q={latitude},{longitude}"

        sites.append(
            {
                "id": str(int(serial)),
                "siteCode": site_code,
                "displayType": clean_str(row[4].value),
                "zone": clean_str(row[2].value),
                "location": clean_str(row[5].value),
                "latitude": latitude,
                "longitude": longitude,
                "width": width,
                "height": height,
                "qty": qty,
                "area": area,
                "mediaStatus": clean_str(row[10].value),
                "litStatus": clean_str(row[11].value),
                "googleMapsUrl": google_maps_url,
            }
        )

    return {
        "generatedFrom": xlsx_path.name,
        "siteCount": len(sites),
        "sites": sites,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx", type=Path)
    parser.add_argument(
        "--out",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "src" / "data" / "sites.json",
    )
    args = parser.parse_args()
    payload = convert(args.xlsx)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {payload['siteCount']} sites to {args.out}")


if __name__ == "__main__":
    main()
